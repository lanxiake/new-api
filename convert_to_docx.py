import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from markdown import markdown
from io import StringIO

# 定义源文件夹和目标文件夹
source_base = r"d:\my-project\new-api\docs\marketing"
target_base = r"d:\my-project\new-api\营销策略文档"

# 定义文档映射关系：源文件 -> 中文名称
doc_mapping = {
    # 01-strategy 策略规划
    "01-strategy/01-overall-strategy-20260520.md": "01-整体营销策略（自然渗透+SEO+社区）.docx",
    "01-strategy/02-paid-promotion-plan-20260520.md": "02-30天付费投流方案（2000元预算）.docx",
    "01-strategy/03-pricing-promotion-playbook-20260520.md": "03-定价与充值促销打法（充值5折、折算话术、投流对齐）.docx",
    # 02-content 内容素材
    "02-content/01-articles-library-20260520.md": "04-长文内容库（技术博客、教程）.docx",
    "02-content/02-social-copy-library-20260520.md": "05-社交媒体文案库（30+平台话术）.docx",
    "02-content/03-creatives-paid-library-20260520.md": "06-付费投放创意素材（小红书、公众号模板）.docx",
    # 03-channels 渠道管理
    "03-channels/01-channels-list-20260520.md": "07-30+推广渠道清单与优先级.docx",
    "03-channels/02-kol-outreach-sop-20260520.md": "08-KOC-KOL合作标准流程.docx",
    # 05-community 社群运营
    "05-community/01-private-traffic-ops-20260520.md": "09-QQ群+微信群双轨运营手册.docx",
    "05-community/02-campaigns-library-20260520.md": "10-活动方案库（邀请奖励、周活动）.docx",
}

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_table_to_doc(doc, table_element):
    """将HTML表格转换为Word表格"""
    rows = table_element.find_all('tr')
    if not rows:
        return

    # 创建Word表格
    num_cols = len(rows[0].find_all(['th', 'td']))
    doc_table = doc.add_table(rows=len(rows), cols=num_cols)
    doc_table.style = 'Table Grid'

    for row_idx, tr in enumerate(rows):
        cells = tr.find_all(['th', 'td'])
        for col_idx, cell in enumerate(cells):
            # 获取单元格内容
            cell_text = cell.get_text(strip=True)

            # 写入Word表格
            doc_cell = doc_table.rows[row_idx].cells[col_idx]
            doc_cell.text = cell_text

            # 如果是表头行，设置加粗和背景色
            if row_idx == 0:
                for paragraph in doc_cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                set_cell_shading(doc_cell, 'D9E2F3')  # 浅蓝色背景

    return doc_table

def parse_html_content(html_content, doc):
    """解析HTML内容并添加到Word文档"""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_content, 'html.parser')

    for element in soup.children:
        if element.name is None:
            continue

        if element.name == 'h1':
            text = element.get_text(strip=True)
            if text:
                doc.add_heading(text, level=1)
        elif element.name == 'h2':
            text = element.get_text(strip=True)
            if text:
                doc.add_heading(text, level=2)
        elif element.name == 'h3':
            text = element.get_text(strip=True)
            if text:
                doc.add_heading(text, level=3)
        elif element.name == 'h4':
            text = element.get_text(strip=True)
            if text:
                doc.add_heading(text, level=4)
        elif element.name == 'p':
            text = element.get_text(strip=True)
            if text:
                doc.add_paragraph(text)
        elif element.name == 'ul':
            for li in element.find_all('li', recursive=False):
                text = li.get_text(strip=True)
                if text:
                    doc.add_paragraph(text, style='List Bullet')
        elif element.name == 'ol':
            for li in element.find_all('li', recursive=False):
                text = li.get_text(strip=True)
                if text:
                    doc.add_paragraph(text, style='List Number')
        elif element.name == 'blockquote':
            text = element.get_text(strip=True)
            if text:
                p = doc.add_paragraph(text)
                p_format = p.paragraph_format
                p_format.left_indent = Inches(0.5)
        elif element.name == 'pre':
            text = element.get_text(strip=True)
            if text:
                p = doc.add_paragraph(text)
                for run in p.runs:
                    run.font.name = 'Courier New'
        elif element.name == 'table':
            add_table_to_doc(doc, element)
        elif element.name == 'strong' or element.name == 'b':
            text = element.get_text(strip=True)
            if text:
                p = doc.add_paragraph(text)
                for run in p.runs:
                    run.bold = True
        elif element.name == 'code':
            text = element.get_text(strip=True)
            if text:
                p = doc.add_paragraph(text)
                for run in p.runs:
                    run.font.name = 'Courier New'
        elif element.name == 'hr':
            doc.add_paragraph('─' * 50)

def convert_markdown_to_docx(md_file_path, docx_file_path):
    """
    将Markdown文件转换为DOCX格式（保留表格样式）
    """
    print(f"正在转换: {md_file_path}")
    print(f"  -> {docx_file_path}")

    # 创建新的Word文档
    doc = Document()

    # 读取Markdown文件内容
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_content = f.read()

    # 将Markdown转换为HTML（启用表格扩展）
    html_content = markdown(md_content, extensions=['tables', 'fenced_code', 'codehilite'])

    # 解析HTML并添加到Word文档
    parse_html_content(html_content, doc)

    # 保存文档
    doc.save(docx_file_path)
    print(f"  完成!")

def main():
    print("=" * 60)
    print("Markdown转DOCX转换工具（增强版-支持表格样式）")
    print("=" * 60)

    # 确保目标文件夹存在
    os.makedirs(target_base, exist_ok=True)

    # 转换每个文档
    for md_path, docx_name in doc_mapping.items():
        full_md_path = os.path.join(source_base, md_path)
        full_docx_path = os.path.join(target_base, docx_name)

        if os.path.exists(full_md_path):
            convert_markdown_to_docx(full_md_path, full_docx_path)
        else:
            print(f"警告: 文件不存在 {full_md_path}")

    print("\n" + "=" * 60)
    print("转换完成!")
    print(f"输出目录: {target_base}")
    print("=" * 60)

if __name__ == "__main__":
    main()